"""Exportaciones: Excel, PDF, comprobante de contrato y respaldo de la base de datos."""
import os
import shutil
import subprocess
from datetime import datetime, date

from database import DB_CONFIG, get_connection
from . import logic
from .services import BACKUPS_DIR


# ------------------------------------------------------------------- Excel
def export_xlsx(path, title, headers, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Falta la librería openpyxl. Ejecutá INSTALAR_DEPENDENCIAS.bat.") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="155BD5")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([("" if v is None else v) for v in r])
    for c in range(1, len(headers) + 1):
        width = max(len(str(headers[c - 1])), *(len(str(r[c - 1])) for r in rows)) if rows else len(str(headers[c - 1]))
        ws.column_dimensions[get_column_letter(c)].width = min(48, width + 3)
    ws.freeze_panes = "A2"
    wb.save(path)


# --------------------------------------------------------------------- PDF
def export_pdf_table(path, title, headers, rows, subtitle=""):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
    except ImportError as exc:
        raise RuntimeError("Falta la librería reportlab. Ejecutá INSTALAR_DEPENDENCIAS.bat.") from exc
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=8, leading=10)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    doc = SimpleDocTemplate(path, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm, title=title)
    story = [Paragraph(f"<b>Foxy Lancer · {title}</b>", styles["Title"])]
    meta = f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} · {len(rows)} registro(s)"
    if subtitle:
        meta += f" · {subtitle}"
    story += [Paragraph(meta, styles["Normal"]), Spacer(1, 10)]
    data = [[Paragraph(str(h), head) for h in headers]]
    for r in rows:
        data.append([Paragraph("" if v is None else str(v).replace("&", "&amp;").replace("<", "&lt;"), cell) for v in r])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#155bd5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef3fa")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c4d0e2")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    doc.build(story)


def contract_pdf(path, c):
    """Comprobante simple del contrato. No es un documento legal ni un comprobante de pago."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
    except ImportError as exc:
        raise RuntimeError("Falta la librería reportlab. Ejecutá INSTALAR_DEPENDENCIAS.bat.") from exc
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                            title=f"Comprobante de trabajo #{c['id']}")
    esc = lambda s: str(s or "").replace("&", "&amp;").replace("<", "&lt;")
    story = [Paragraph("<b>Foxy Lancer · Comprobante de trabajo</b>", styles["Title"]),
             Paragraph(f"Contrato #{c['id']} · emitido el {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]), Spacer(1, 14)]
    info = [["Trabajo", esc(c["titulo"])], ["Publicador", esc(c["publicador"])],
            ["Candidato", f"{esc(c['candidato'])} ({esc(c['candidato_correo'])})"],
            ["Monto acordado", logic.money(c["monto_acordado"])], ["Fecha de entrega", logic.fmt_dt(c["fecha_entrega"]) or "—"],
            ["Estado", c["estado"]]]
    t = Table(info, colWidths=[4.5 * cm, 12 * cm])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef3fa")),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c4d0e2")),
                           ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"), ("PADDING", (0, 0), (-1, -1), 6)]))
    story += [t, Spacer(1, 16), Paragraph("<b>Hitos</b>", styles["Heading2"])]
    rows = [["Hito", "Fecha límite", "Estado"]]
    for h in c.get("hitos", []):
        rows.append([esc(h["titulo"]), logic.fmt_dt(h["fecha_limite"]) or "—", "Completado" if h["completado"] else "Pendiente"])
    if len(rows) == 1:
        rows.append(["Sin hitos registrados", "—", "—"])
    ht = Table(rows, colWidths=[9 * cm, 3.5 * cm, 4 * cm], repeatRows=1)
    ht.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#155bd5")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c4d0e2")), ("PADDING", (0, 0), (-1, -1), 5)]))
    story += [ht, Spacer(1, 18),
              Paragraph("<i>Este comprobante resume el acuerdo registrado en Foxy Lancer. No es un documento legal ni un comprobante de pago.</i>",
                        styles["Normal"])]
    doc.build(story)


# ---------------------------------------------------------------- respaldo
def _sql_value(v):
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, (bytes, bytearray)):
        return "0x" + bytes(v).hex() if v else "''"
    if isinstance(v, (datetime, date)):
        return "'" + v.isoformat(sep=" ") + "'" if isinstance(v, datetime) else "'" + v.isoformat() + "'"
    s = str(v)
    for a, b in (("\\", "\\\\"), ("'", "\\'"), ("\n", "\\n"), ("\r", "\\r"), ("\x00", "\\0"), ("\x1a", "\\Z")):
        s = s.replace(a, b)
    return "'" + s + "'"


def backup_database():
    """Crea un respaldo .sql en backups/. Solo lee la base de datos."""
    os.makedirs(BACKUPS_DIR, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(BACKUPS_DIR, f"mypines_{stamp}.sql")
    dump = shutil.which("mysqldump")
    if dump:
        env = dict(os.environ, MYSQL_PWD=str(DB_CONFIG["password"]))
        cmd = [dump, f"--host={DB_CONFIG['host']}", f"--port={DB_CONFIG['port']}", f"--user={DB_CONFIG['user']}",
               "--single-transaction", "--routines", "--default-character-set=utf8mb4", f"--result-file={path}", DB_CONFIG["database"]]
        try:
            result = subprocess.run(cmd, env=env, capture_output=True, text=True, timeout=120)
            if result.returncode == 0 and os.path.exists(path) and os.path.getsize(path) > 0:
                return path
        except Exception:
            pass
        if os.path.exists(path):
            os.remove(path)
    cn = get_connection()
    try:
        cur = cn.cursor()
        cur.execute("SHOW FULL TABLES WHERE Table_type='BASE TABLE'")
        tables = [r[0] for r in cur.fetchall()]
        with open(path, "w", encoding="utf-8") as f:
            f.write(f"-- Respaldo Foxy Lancer generado el {datetime.now():%Y-%m-%d %H:%M:%S}\n")
            f.write("SET FOREIGN_KEY_CHECKS=0;\nSET NAMES utf8mb4;\n\n")
            for table in tables:
                cur.execute(f"SHOW CREATE TABLE `{table}`")
                ddl = cur.fetchone()[1]
                f.write(f"DROP TABLE IF EXISTS `{table}`;\n{ddl};\n\n")
                cur.execute(f"SELECT * FROM `{table}`")
                cols = [d[0] for d in cur.description]
                col_sql = ", ".join(f"`{c}`" for c in cols)
                while True:
                    batch = cur.fetchmany(200)
                    if not batch:
                        break
                    values = ",\n".join("(" + ", ".join(_sql_value(v) for v in row) + ")" for row in batch)
                    f.write(f"INSERT INTO `{table}` ({col_sql}) VALUES\n{values};\n")
                f.write("\n")
            f.write("SET FOREIGN_KEY_CHECKS=1;\n")
    finally:
        cn.close()
    return path


def list_backups():
    if not os.path.isdir(BACKUPS_DIR):
        return []
    out = []
    for name in sorted(os.listdir(BACKUPS_DIR), reverse=True):
        if name.endswith(".sql"):
            p = os.path.join(BACKUPS_DIR, name)
            out.append({"nombre": name, "ruta": p, "tamano": os.path.getsize(p), "fecha": datetime.fromtimestamp(os.path.getmtime(p))})
    return out
